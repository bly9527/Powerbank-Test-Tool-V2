import openpyxl
from datetime import time, datetime

def get_target_soc_list(mode='discharge'):
    """
    获取需要填写的 28 个目标 SOC 百分比列表
    """
    if mode == 'discharge':
        # 放电: 99~90 (10个), 80~10 (8个), 9~0 (10个)
        front = list(range(99, 89, -1))
        mid = [80, 70, 60, 50, 40, 30, 20, 10]
        back = list(range(9, -1, -1))
        return front + mid + back
    else:
        # 充电: 1~10 (10个), 20~90 (8个), 91~100 (10个)
        front = list(range(1, 11))
        mid = [20, 30, 40, 50, 60, 70, 80, 90]
        back = list(range(91, 101))
        return front + mid + back

def get_cell_coords(soc, mode='discharge', base_row=0, col_offset=0):
    """
    计算特定 SOC 对应的 (time_row, time_col) 和 (cap_row, cap_col)
    支持 col_offset: 1# 表为 0, 2# 表为 14
    """
    if mode == 'discharge':
        if 90 <= soc <= 99:
            # 99 对应 col 2, 98 对应 col 3 ... 90 对应 col 11
            col = 2 + (99 - soc) + col_offset
            time_row = base_row + 4
            cap_row = base_row + 5
            return (time_row, col), (cap_row, col)
        elif soc in [80, 70, 60, 50, 40, 30, 20, 10]:
            idx = (80 - soc) // 10  # 0, 1, ..., 7
            time_row = base_row + 7 + idx * 3
            cap_row = base_row + 8 + idx * 3
            col = 11 + col_offset
            return (time_row, col), (cap_row, col)
        elif 0 <= soc <= 9:
            col = 2 + (9 - soc) + col_offset
            time_row = base_row + 31
            cap_row = base_row + 32
            return (time_row, col), (cap_row, col)
    else:  # charge
        if 1 <= soc <= 10:
            col = 2 + (soc - 1) + col_offset
            time_row = base_row + 4
            cap_row = base_row + 5
            return (time_row, col), (cap_row, col)
        elif soc in [20, 30, 40, 50, 60, 70, 80, 90]:
            idx = (soc - 20) // 10  # 0, 1, ..., 7
            time_row = base_row + 7 + idx * 3
            cap_row = base_row + 8 + idx * 3
            col = 11 + col_offset
            return (time_row, col), (cap_row, col)
        elif 91 <= soc <= 100:
            col = 2 + (soc - 91) + col_offset
            time_row = base_row + 31
            cap_row = base_row + 32
            return (time_row, col), (cap_row, col)

    return None, None

def clean_sheet_data_and_images(ws):
    """
    清除工作表中的所有图片和历史测试数据，100% 保留所有公式与表头格式
    """
    # 1. 删除所有图片
    if hasattr(ws, '_images'):
        ws._images.clear()

    # 2. 定位所有的测试数据块
    title_rows = set()
    for r in range(1, ws.max_row + 1):
        v1 = ws.cell(r, 1).value
        v2 = ws.cell(r, 15).value
        for v in (v1, v2):
            if v and any(k in str(v) for k in ['PB', '放电', '充电', '数显测试']):
                base_r = r - 2 if r >= 2 else 0
                title_rows.add(base_r)

    # 3. 清理每个测试块中 1# 和 2# 的测试数据
    for base_row in title_rows:
        for col_offset in [0, 14]:
            # 顶部 10 个数据点 (base_row+4 为时间, base_row+5 为容量)
            for c in range(2 + col_offset, 12 + col_offset):
                cell_t = ws.cell(base_row + 4, c)
                if cell_t.value and not str(cell_t.value).startswith('='):
                    cell_t.value = None
                cell_c = ws.cell(base_row + 5, c)
                if cell_c.value and not str(cell_c.value).startswith('='):
                    cell_c.value = None

            # 中间 8 个跨度数据点 (第 11 列 / 25 列)
            for idx in range(8):
                cell_t = ws.cell(base_row + 7 + idx * 3, 11 + col_offset)
                if cell_t.value and not str(cell_t.value).startswith('='):
                    cell_t.value = None
                cell_c = ws.cell(base_row + 8 + idx * 3, 11 + col_offset)
                if cell_c.value and not str(cell_c.value).startswith('='):
                    cell_c.value = None

            # 底部 10 个数据点 (base_row+31 为时间, base_row+32 为容量)
            for c in range(2 + col_offset, 12 + col_offset):
                cell_t = ws.cell(base_row + 31, c)
                if cell_t.value and not str(cell_t.value).startswith('='):
                    cell_t.value = None
                cell_c = ws.cell(base_row + 32, c)
                if cell_c.value and not str(cell_c.value).startswith('='):
                    cell_c.value = None

            # 统计汇总行 (base_row+33) 的历史值
            for c in [2 + col_offset, 7 + col_offset, 11 + col_offset]:
                cell_sum = ws.cell(base_row + 33, c)
                if cell_sum.value and not str(cell_sum.value).startswith('='):
                    cell_sum.value = None

def scan_excel_test_groups(excel_path):
    """
    扫描 Excel 模板，获取所有可用的测试组
    返回: dict {sheet_name: [{'name': title, 'base_row': int, 'title_row': int, 'col_offset': int}]}
    """
    groups = {}
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_groups = []
        for r in range(1, ws.max_row + 1):
            # 扫描 1# (列 1)
            val1 = ws.cell(r, 1).value
            if val1 and ('PB' in str(val1) or '放电' in str(val1) or '充电' in str(val1)):
                base_row = r - 2 if r >= 2 else 0
                sheet_groups.append({
                    'name': str(val1).strip(),
                    'base_row': base_row,
                    'title_row': r,
                    'sheet': sheet_name,
                    'col_offset': 0
                })

            # 扫描 2# (列 15)
            val2 = ws.cell(r, 15).value
            if val2 and ('PB' in str(val2) or '放电' in str(val2) or '充电' in str(val2)):
                base_row = r - 2 if r >= 2 else 0
                sheet_groups.append({
                    'name': str(val2).strip(),
                    'base_row': base_row,
                    'title_row': r,
                    'sheet': sheet_name,
                    'col_offset': 14
                })

        if sheet_groups:
            groups[sheet_name] = sheet_groups
    return groups

def fill_excel_report(excel_path, sheet_name, base_row, mode,
                      extracted_data=None, output_path=None, col_offset=0,
                      extracted_data_2=None):
    """
    将提取/修改的数据写入 Excel 模板并保存。
    特点:
    1. 自动删除所有图片与历史残留数据；
    2. 仅保留当前测试类型的工作表，移除其它无关表格；
    3. 支持单台样机写入，也支持 1# 和 2# 双台样机同时回填；
    4. 完整保留原表的全部公式和计算逻辑。
    """
    wb = openpyxl.load_workbook(excel_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"工作表 '{sheet_name}' 不存在！")

    # 导出时只保留当前测试的一个模板表，删除其它工作表
    for sname in list(wb.sheetnames):
        if sname != sheet_name:
            wb.remove(wb[sname])

    ws = wb[sheet_name]

    # 清空图片与旧数据
    clean_sheet_data_and_images(ws)

    filled_count_1 = 0
    filled_count_2 = 0
    target_socs = get_target_soc_list(mode)

    # 写入 1# 数据
    if extracted_data:
        for soc in target_socs:
            if soc in extracted_data:
                item = extracted_data[soc]
                time_val = item.get('time_obj')
                cap_val = item.get('capacity')

                time_pos, cap_pos = get_cell_coords(soc, mode=mode, base_row=base_row, col_offset=col_offset)
                if time_pos and time_val is not None:
                    t_cell = ws.cell(time_pos[0], time_pos[1])
                    t_cell.value = time_val
                    t_cell.number_format = 'h:mm:ss'

                if cap_pos and cap_val is not None:
                    c_cell = ws.cell(cap_pos[0], cap_pos[1])
                    c_cell.value = float(cap_val)
                    c_cell.number_format = '0.0000_ '

                filled_count_1 += 1

    # 写入 2# 数据 (固定在第 15 列起算，对应 col_offset = 14)
    if extracted_data_2:
        for soc in target_socs:
            if soc in extracted_data_2:
                item = extracted_data_2[soc]
                time_val = item.get('time_obj')
                cap_val = item.get('capacity')

                time_pos, cap_pos = get_cell_coords(soc, mode=mode, base_row=base_row, col_offset=14)
                if time_pos and time_val is not None:
                    t_cell = ws.cell(time_pos[0], time_pos[1])
                    t_cell.value = time_val
                    t_cell.number_format = 'h:mm:ss'

                if cap_pos and cap_val is not None:
                    c_cell = ws.cell(cap_pos[0], cap_pos[1])
                    c_cell.value = float(cap_val)
                    c_cell.number_format = '0.0000_ '

                filled_count_2 += 1

    save_target = output_path if output_path else excel_path
    wb.save(save_target)
    return filled_count_1, filled_count_2
