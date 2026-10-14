"""Persistent test projects and image-preserving Excel report updates."""
import json
import math
import os
import shutil
import tempfile
from copy import copy, deepcopy
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from core.excel_filler import get_cell_coords, get_target_soc_list, scan_excel_test_groups


def projects_home():
    return Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / 'PowerbankTestTool' / 'projects'


def group_key(group):
    return f"{group['sheet']}|{group['title_row']}|{group['col_offset']}"


def project_catalog():
    root = projects_home()
    if not root.exists():
        return []
    result = []
    for path in root.iterdir():
        metadata = path / 'project.json'
        if metadata.is_file():
            try:
                info = json.loads(metadata.read_text(encoding='utf-8'))
                result.append({'name': info['name'], 'path': str(path), 'updated_at': info.get('updated_at', '')})
            except (OSError, ValueError, KeyError):
                continue
    return sorted(result, key=lambda entry: entry['updated_at'], reverse=True)


def delete_project(root):
    """Delete one registered project directory after validating its boundary."""
    home = projects_home().resolve()
    root = Path(root).resolve()
    if root.parent != home or not root.name.startswith('project_'):
        raise ValueError('只能删除软件项目目录中的存档')
    project = ProjectArchive(root)
    project._ensure_excel_closed()
    shutil.rmtree(project.root)


def _atomic_json(path, value):
    path = Path(path)
    temporary = path.with_name(path.name + '.' + uuid4().hex + '.tmp')
    try:
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def create_project(name, template_path):
    name = name.strip()
    source = Path(template_path)
    if not name or not source.is_file():
        raise ValueError('请输入存档名称并选择有效的 Excel 模板')
    if not scan_excel_test_groups(str(source)):
        raise ValueError('模板中没有可识别的充电/放电测试项')
    root = projects_home() / ('project_' + uuid4().hex)
    root.mkdir(parents=True)
    try:
        shutil.copyfile(source, root / 'report.xlsx')
        # The supplied PB519 sheet has ten 2# charge percentage formulas in
        # its first block referencing the 1# total ($K$32). That fails when
        # only 2# is tested; use its own total in the working copy.
        workbook = openpyxl.load_workbook(root / 'report.xlsx')
        charge = workbook['充电时间及曲线'] if '充电时间及曲线' in workbook else None
        if charge:
            for row in range(5, 33, 3):
                cell = charge[f'AA{row}']
                if cell.value == f'=Z{row}/$K$32':
                    cell.value = f'=Z{row}/$Y$32'
        workbook.save(root / 'report.xlsx')
        workbook.close()
        now = datetime.now().isoformat(timespec='seconds')
        _atomic_json(root / 'project.json', {'version': 1, 'name': name,
            'template_name': source.name, 'created_at': now, 'updated_at': now,
            'entries': {}, 'planned': []})
    except Exception:
        shutil.rmtree(root)
        raise
    return ProjectArchive(root)


class ProjectArchive:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.metadata_path = self.root / 'project.json'
        self.report_path = self.root / 'report.xlsx'
        self.metadata = json.loads(self.metadata_path.read_text(encoding='utf-8'))
        if not self.report_path.is_file():
            raise FileNotFoundError(f'存档工作版 Excel 不存在：{self.report_path}')

    @property
    def groups(self):
        return scan_excel_test_groups(str(self.report_path))

    def _ensure_excel_closed(self):
        if (self.root / ('~$' + self.report_path.name)).exists():
            raise PermissionError('工作版 Excel 正在打开。请先在 Excel 中保存并关闭，再写入或导出。')

    def _evidence_path(self, value):
        """Resolve archived evidence without following paths outside this project."""
        if not isinstance(value, str) or not value or '\x00' in value:
            return None
        try:
            path = Path(value)
            if '..' in path.parts:
                return None
            evidence = (self.root / 'evidence').resolve()
            if evidence.parent != self.root:
                return None
            resolved = (path if path.is_absolute() else self.root / path).resolve()
            if resolved.parent != evidence or resolved.suffix.lower() != '.png':
                return None
            return resolved if resolved.is_file() else None
        except (OSError, ValueError):
            return None

    def resolve_record(self, record, source_video=''):
        """Return a UI copy with a safe absolute evidence path; legacy rows still work."""
        result = dict(record)
        path = self._evidence_path(result.get('snapshot_path'))
        result['snapshot_path'] = str(path) if path else ''
        source = result.get('source_video')
        result['source_video'] = source if isinstance(source, str) and source else str(source_video or '')
        return result

    def _archive_record(self, record, source_video):
        result = {'time_str': record.get('time_str', ''),
                  'capacity': record.get('capacity'), 'video_sec': record.get('video_sec')}
        result.update({key: record[key] for key in ('capture_method', 'capture_unit', 'manual_override', 'capacity_raw',
                  'timing_method', 'timing_original_sec', 'timing_search_start_sec') if key in record})
        source = record.get('source_video')
        result['source_video'] = source if isinstance(source, str) and source else str(source_video or '')
        if record.get('locator_version') != 1:
            return result
        seconds = record.get('video_sec')
        index = record.get('frame_index')
        if (isinstance(seconds, bool) or not isinstance(seconds, (int, float))
                or not math.isfinite(seconds) or seconds < 0
                or isinstance(index, bool) or not isinstance(index, int) or index < 0):
            return result
        result.update(locator_version=1, frame_index=index)
        snapshot = record.get('snapshot_path')
        existing = self._evidence_path(snapshot)
        if existing:
            result['snapshot_path'] = existing.relative_to(self.root).as_posix()
            return result
        # Live worker snapshots are absolute. Never interpret an arbitrary relative
        # path against the process working directory or escape the project root.
        if not isinstance(snapshot, str) or not snapshot or '\x00' in snapshot:
            return result
        try:
            source_path = Path(snapshot)
            if not source_path.is_absolute() or source_path.suffix.lower() != '.png' or not source_path.is_file():
                return result
            evidence = self.root / 'evidence'
            if evidence.resolve().parent != self.root:
                raise ValueError('存档截图目录不在当前存档内')
            evidence.mkdir(exist_ok=True)
            destination = evidence / (uuid4().hex + '.png')
            temporary = destination.with_suffix('.tmp')
            try:
                shutil.copyfile(source_path, temporary)
                os.replace(temporary, destination)
            finally:
                temporary.unlink(missing_ok=True)
            result['snapshot_path'] = destination.relative_to(self.root).as_posix()
        except (FileNotFoundError, NotADirectoryError):
            # The temporary live image may already have been cleaned up. Do not
            # persist a broken evidence reference, but keep usable video location.
            pass
        return result

    def _validate_group(self, group, unit, mode):
        if mode not in ('charge', 'discharge') or unit not in (1, 2):
            raise ValueError('无效的测试类型或样机编号')
        sheet = '充电时间及曲线' if mode == 'charge' else '放电时间及曲线'
        expected_offset = 0 if unit == 1 else 14
        if group.get('sheet') != sheet or group.get('col_offset') != expected_offset:
            raise ValueError(f'{unit}# 选择的 Excel 测试项与充放电类型不对应')
        if not any(group_key(item) == group_key(group) for item in self.groups.get(sheet, [])):
            raise ValueError('存档工作版 Excel 中找不到此测试项')

    @staticmethod
    def _clear_and_write(ws, group, mode, records):
        count = 0
        for soc in get_target_soc_list(mode):
            time_cell, cap_cell = get_cell_coords(soc, mode, group['base_row'], group['col_offset'])
            for pos in (time_cell, cap_cell):
                cell = ws.cell(*pos)
                if cell.data_type != 'f':
                    cell.value = None
            record = records.get(soc)
            if record is None:
                continue
            time_value = record.get('time_obj')
            if time_value is None and record.get('time_str'):
                try:
                    time_value = datetime.strptime(record['time_str'], '%H:%M:%S').time()
                except ValueError:
                    pass
            if time_value is not None:
                ws.cell(*time_cell).value = time_value
                ws.cell(*time_cell).number_format = 'h:mm:ss'
            capacity = record.get('capacity')
            if capacity is not None:
                ws.cell(*cap_cell).value = float(capacity)
                ws.cell(*cap_cell).number_format = '0.0000_ '
            count += 1
        return count

    def save_results(self, assignments, video_path):
        """Update only the selected blocks; preserve pictures and other blocks."""
        self._ensure_excel_closed()
        if not assignments:
            raise ValueError('没有要存入项目的数据')
        for item in assignments:
            self._validate_group(item['group'], item['unit'], item['mode'])
            if not item['records']:
                raise ValueError(f"{item['unit']}# 没有可保存的数据")
        # Finish evidence copies before opening or replacing the workbook. A
        # failed copy must not publish metadata pointing at an absent image.
        archived_records = [{str(soc): self._archive_record(rec, video_path)
                             for soc, rec in item['records'].items()}
                            for item in assignments]
        original_mtime = self.report_path.stat().st_mtime_ns
        workbook = openpyxl.load_workbook(self.report_path)
        metadata = json.loads(json.dumps(self.metadata))
        counts = {}
        now = datetime.now().isoformat(timespec='seconds')
        for item, stored_records in zip(assignments, archived_records):
            group = item['group']
            key = group_key(group)
            count = self._clear_and_write(workbook[group['sheet']], group, item['mode'], item['records'])
            counts[item['unit']] = count
            metadata['entries'][key] = {'unit': item['unit'], 'mode': item['mode'],
                'group': group, 'video_path': str(video_path or ''), 'updated_at': now,
                'records': stored_records}
        metadata['updated_at'] = now
        fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=self.root)
        os.close(fd)
        try:
            workbook.save(temporary)
            if self.report_path.stat().st_mtime_ns != original_mtime:
                raise IOError('Excel 文件在写入期间发生变化，请重新打开存档后再试。')
            os.replace(temporary, self.report_path)
            _atomic_json(self.metadata_path, metadata)
        finally:
            workbook.close()
            Path(temporary).unlink(missing_ok=True)
        self.metadata = metadata
        return counts

    def set_planned(self, keys):
        valid = {group_key(group) for groups in self.groups.values() for group in groups}
        self.metadata['planned'] = sorted(set(keys) & valid)
        self.metadata['updated_at'] = datetime.now().isoformat(timespec='seconds')
        _atomic_json(self.metadata_path, self.metadata)

    def export(self, output_path):
        """Create a compact report with only completed blocks and their pictures."""
        self._ensure_excel_closed()
        output_path = Path(output_path).resolve()
        if output_path == self.report_path:
            raise ValueError('导出路径不能覆盖存档的工作版 Excel')
        selected = set(self.metadata['entries'])
        if not selected:
            raise ValueError('存档还没有保存任何测试结果')
        source = openpyxl.load_workbook(self.report_path)
        result = openpyxl.Workbook()
        result.remove(result.active)
        groups = self.groups
        for sheet_name in source.sheetnames:
            all_groups = groups.get(sheet_name, [])
            active = [g for g in all_groups if group_key(g) in selected]
            if not active:
                continue
            ws = source[sheet_name]
            target = result.create_sheet(sheet_name)
            for attr in ('sheet_format', 'sheet_properties', 'page_setup',
                         'page_margins', 'print_options'):
                setattr(target, attr, copy(getattr(ws, attr)))
            for key, dimension in ws.column_dimensions.items():
                dest_dimension = target.column_dimensions[key]
                dest_dimension.width = dimension.width
                dest_dimension.hidden = dimension.hidden
                dest_dimension.bestFit = dimension.bestFit
                dest_dimension.outlineLevel = dimension.outlineLevel
            title_rows = sorted({g['title_row'] for g in all_groups})
            output_start = 1
            output_max_col = 1
            for index, title_row in enumerate(title_rows):
                chosen = [g for g in active if g['title_row'] == title_row]
                if not chosen:
                    continue
                start = max(1, title_row - 1)
                end = title_rows[index + 1] - 2 if index + 1 < len(title_rows) else ws.max_row
                shift = output_start - start
                allowed = set()
                for group in chosen:
                    first = 1 if group['col_offset'] == 0 else 15
                    allowed.update(range(first, first + 13))
                output_max_col = max(output_max_col, max(allowed))
                for source_row in range(start, end + 1):
                    dest_row = source_row + shift
                    if source_row in ws.row_dimensions:
                        dimension = ws.row_dimensions[source_row]
                        target.row_dimensions[dest_row].height = dimension.height
                        target.row_dimensions[dest_row].hidden = dimension.hidden
                        target.row_dimensions[dest_row].outlineLevel = dimension.outlineLevel
                    for col in allowed:
                        current = ws.cell(source_row, col)
                        destination = target.cell(dest_row, col)
                        if current.data_type == 'f':
                            destination.value = Translator(current.value, origin=current.coordinate).translate_formula(
                                destination.coordinate)
                        elif current.value is not None:
                            destination.value = current.value
                        if current.has_style:
                            destination.font = copy(current.font)
                            destination.fill = copy(current.fill)
                            destination.border = copy(current.border)
                            destination.alignment = copy(current.alignment)
                            destination.protection = copy(current.protection)
                            destination.number_format = current.number_format
                        if current.hyperlink:
                            destination._hyperlink = copy(current.hyperlink)
                        if current.comment:
                            destination.comment = copy(current.comment)
                for merged in ws.merged_cells.ranges:
                    if (merged.min_row >= start and merged.max_row <= end
                            and all(col in allowed for col in range(merged.min_col, merged.max_col + 1))):
                        target.merge_cells(start_row=merged.min_row + shift,
                            end_row=merged.max_row + shift,
                            start_column=merged.min_col, end_column=merged.max_col)
                for image in ws._images:
                    anchor = image.anchor
                    image_row = anchor._from.row + 1 if hasattr(anchor, '_from') else None
                    if image_row is None or not start <= image_row <= end:
                        continue
                    copied = deepcopy(image)
                    copied.anchor._from.row += shift
                    if hasattr(copied.anchor, 'to'):
                        copied.anchor.to.row += shift
                    target.add_image(copied)
                output_start += end - start + 1
            target.print_area = f'A1:{get_column_letter(output_max_col)}{output_start - 1}'
        fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=output_path.parent)
        os.close(fd)
        try:
            result.save(temporary)
            os.replace(temporary, output_path)
        finally:
            source.close()
            result.close()
            Path(temporary).unlink(missing_ok=True)
        return len(selected)
