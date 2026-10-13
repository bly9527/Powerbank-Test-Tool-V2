from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'release' / '充电宝数显均匀度测试工具使用攻略.docx'
BLUE = '1F4E78'
PALE_BLUE = 'EAF2F8'
LIGHT_GRAY = 'F5F6F7'
BORDER = 'D9D9D9'


def set_font(run, name='Microsoft YaHei', size=10.5, bold=False, color='000000'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_borders(cell, color=BORDER, size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), size)
        node.set(qn('w:color'), color)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tc_pr.append(margins)
    for name, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = margins.find(qn('w:' + name))
        if node is None:
            node = OxmlElement('w:' + name)
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement('w:tblHeader')
    header_flag.set(qn('w:val'), 'true')
    tr_pr.append(header_flag)
    for idx, text in enumerate(headers):
        cell = t.rows[0].cells[idx]
        shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=9.5, bold=True, color='FFFFFF')
    for row_idx, values in enumerate(rows):
        cells = t.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            if row_idx % 2:
                shade(cell, LIGHT_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)), size=9.2)
    for row in t.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_margin(cell)
            if widths:
                cell.width = Cm(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text))
    return p


_number = 0


def reset_numbers():
    global _number
    _number = 0


def numbered(doc, text):
    global _number
    _number += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(f'{_number}. {text}'))
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('第 ')
    set_font(run, size=9, color='666666')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)
    run = paragraph.add_run(' 页')
    set_font(run, size=9, color='666666')


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Microsoft YaHei'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal.font.size = Pt(10.5)
for style_name, size, before, after in (
        ('Title', 24, 0, 18), ('Heading 1', 16, 18, 8), ('Heading 2', 12.5, 12, 5)):
    style = styles[style_name]
    style.font.name = 'Microsoft YaHei'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
title_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)

footer = section.footer.paragraphs[0]
set_font(footer.add_run('充电宝数显均匀度测试自动数据提取工具'), size=9, color='666666')
footer.add_run('\t')
add_page_number(footer)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(title.add_run('充电宝数显均匀度测试工具使用攻略'), size=24, bold=True)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(subtitle.add_run('适用于 PowerbankTestTool v2 单文件版'), size=11.5, color='555555')
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.paragraph_format.space_before = Pt(18)
intro.paragraph_format.space_after = Pt(22)
set_font(intro.add_run('本攻略面向第一次使用软件的测试人员，覆盖建档、视频识别、数码管样本库、Excel 保存与最终导出的完整流程。'), size=11)

heading(doc, '开始前先看', 1)
table(doc, ['项目', '说明'], [
    ('运行文件', 'PowerbankTestTool_v2_单文件版.exe，直接双击运行，无需安装 Python。'),
    ('适用系统', 'Windows 10 或 Windows 11。首次启动需要解压内部组件，等待时间会比后续启动长。'),
    ('准备材料', '录屏视频、Excel 原始模板；数码管模式还需要人工标注的数字样本。'),
    ('本机数据', '存档和样本库默认保存在当前电脑的 LocalAppData 中，不会随 EXE 自动发给别人。'),
], widths=[3.2, 12.8])

heading(doc, '最快完成一次测试', 1)
flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
flow.paragraph_format.space_after = Pt(12)
set_font(flow.add_run('新建存档  →  打开视频  →  选择样机和测试项  →  画框  →  测试当前帧  →  自动分析  →  保存到存档  →  导出 Excel'), size=10.5, bold=True, color=BLUE)
reset_numbers()
for text in [
    '新建存档，并选择这次项目使用的 Excel 原始模板。',
    '打开录屏视频，勾选参与测试的 1# 或 2# 样机，分别选择充电或放电以及对应的 Excel 测试项。',
    '在清晰画面上完成电量、时间、容量三个选区，点击“测试当前帧框选识别”。',
    '确认单帧结果合理后开始自动分析，分析结束后核对并修正表格。',
    '点击“保存本次结果到存档”。所有视频完成后导出存档完整 Excel 报告。',
]:
    numbered(doc, text)

heading(doc, '一 建立和管理测试存档', 1)
body(doc, '存档用于把同一型号、同一份 Excel 模板下的多段视频结果连续保存起来。一个视频做完后直接保存到同一存档，不需要每次导出再手工拼表。')
heading(doc, '新建存档', 2)
reset_numbers()
for text in [
    '点击“新建存档”。',
    '输入容易辨认的名称，例如 PB519 9V 测试。',
    '选择 Excel 原始模板。单文件版运行时会释放内置模板，也可以选择你自己的兼容模板。',
    '创建成功后，存档名称会显示在下拉框中，状态区会显示已保存项目数量和工作表路径。',
]: numbered(doc, text)
body(doc, '建议：一个产品型号或一整套测试使用一个存档。1# 和 2# 可以分别选择充电或放电，也可以同时使用不同测试项。', '建议：')

heading(doc, '切换和删除存档', 2)
bullet(doc, '切换存档前，如果当前表格有未保存结果，软件会先询问是否放弃，避免结果串到另一个项目。')
bullet(doc, '视频正在分析时不能切换或删除存档。')
bullet(doc, '删除存档时会显示存档名称、已保存项目数和工作目录，并要求再次确认。')
bullet(doc, '删除会移除该存档的工作版 Excel 和保存记录，不能撤销；最终另存到其他位置的 Excel 不受影响。')
body(doc, '注意：存档 Excel 正在 WPS 或 Excel 中打开时，请先保存并关闭，再执行写入、导出或删除。', '注意：')

heading(doc, '二 打开视频和选择测试项目', 1)
reset_numbers()
for text in [
    '点击“打开录屏视频文件”，选择 MP4、AVI、MKV 或 MOV 视频。',
    '勾选本次真正参与识别的样机。只做一台时取消另一台，避免软件等待不存在的选区。',
    '为每台样机分别选择“充电”或“放电”。两台可以一台充电、一台放电。',
    '选择与样机编号、充放电类型对应的 Excel 测试项。',
]: numbered(doc, text)
body(doc, '鼠标滚轮不会改变左侧数字输入框或下拉选项。需要修改时请点击控件后键盘输入，或使用控件右侧的上下箭头。')

heading(doc, '目标记录点', 2)
table(doc, ['模式', '自动记录的 28 个节点'], [
    ('放电', '99% 到 90%；80%、70%、60%、50%、40%、30%、20%、10%；9% 到 0%。'),
    ('充电', '1% 到 10%；20%、30%、40%、50%、60%、70%、80%、90%；91% 到 100%。'),
], widths=[3, 13])

heading(doc, '三 正确画出识别选区', 1)
body(doc, '先拖动选帧滑块，停在数字完整、清晰且没有切换残影的画面。每台样机需要三个选区。')
table(doc, ['样机', '电量', '时间', '容量'], [
    ('1#', '蓝框', '绿框', '橙框'),
    ('2#', '紫框', '青框', '红框'),
], widths=[3, 4.3, 4.3, 4.3])

heading(doc, '画框原则', 2)
bullet(doc, '电量框只包含完整电量数字，尽量排除百分号、闪电图标、电池边框和其他文字。')
bullet(doc, '时间框覆盖 Power Z 的完整时间。若画面中有“全部”和“窗口”两行，优先让上方“全部”时间清晰可见。')
bullet(doc, '容量框覆盖完整数值和必要的小数点，避免把相邻时间或单位栏一起框入。')
bullet(doc, '设备侧放或倒立时，为对应选区选择旋转角度，使送入识别器的数字恢复正向。')
bullet(doc, '视频位置变化后重新核对选区，不能直接沿用另一段视频的框。')

heading(doc, '识别前检查', 2)
reset_numbers()
numbered(doc, '点击“测试当前帧框选识别”。')
numbered(doc, '查看 1# 和 2# 的电量、时间、容量结果。')
numbered(doc, '结果不正确时，先调整框的位置和大小，再换一帧测试。')
numbered(doc, '单帧结果稳定后再开始自动分析。')

heading(doc, '四 选择识别方式和参数', 1)
table(doc, ['设置', '怎么选'], [
    ('电量屏字形', '普通液晶数字选择“LCD 普通数字”；断码或发光数码管选择“数码管样本库”。'),
    ('抽帧间隔', '数值越小越不容易错过跳变，但分析时间越长。数码管闪烁明显时可从 0.2 秒开始验证。'),
    ('防抖确认帧数', '连续识别到相同数字后才确认，默认值适合先做验证；过高会延迟确认。'),
    ('容量单位', '按视频中上位机实际显示选择 Ah 或 mAh。'),
    ('识别计算线程数', '通常选“自动”。电脑卡顿时可降低，速度不足时可在不过热的前提下增加。'),
], widths=[4, 12])
body(doc, 'CPU 和 GPU 区域显示整机占用情况。GPU 数值只表示显卡当前负载，不代表识别一定使用 GPU。')

heading(doc, '五 普通 LCD 模式操作', 1)
reset_numbers()
for text in [
    '在“电量屏字形”中选择“LCD 普通数字”。',
    '完成电量、时间、容量选区并测试当前帧。',
    '点击“开始自动快进分析”。识别结果会实时进入 1#、2# 数据表。',
    '分析完成后查看捕获点数和状态。点击表格中的视频时刻可以跳回对应画面。',
    '双击时间或容量单元格可以手工修正。确认后保存本次结果到存档。',
]: numbered(doc, text)

heading(doc, '六 数码管样本库模式操作', 1)
body(doc, '数码管模式只改变电量数字的识别方式。Power Z 的时间和容量仍由普通 OCR 读取。不同外观、颜色、亮度或拍摄角度建议建立不同模板。')

heading(doc, '建立模板并采集图片', 2)
reset_numbers()
for text in [
    '选择“数码管样本库”，点击“数码管视频采集 标注 选择样本库”。',
    '点击“新建模板”并命名，例如 PB516 白色数码管。',
    '添加视频，选择采集画面 1# 或 2#，在预览画面中只框住完整电量数字。',
    '设置开始秒、结束秒和采集间隔。闪烁明显时可用 0.2 秒，先采集一小段检查效果。',
    '点击“开始采集当前视频的数字截图”。停止采集会保留已经生成的图片。',
]: numbered(doc, text)

heading(doc, '标注样本', 2)
bullet(doc, '看到完整清晰数字时，输入真实电量并点击“保存标签并查看下一张”。标签立即保存，不需要额外保存按钮。')
bullet(doc, '缺笔画、闪烁、遮挡或无法确定时，选择“不可辨认”，不要根据前后画面猜数字。')
bullet(doc, '不必标完全部图片。优先覆盖需要识别的电量，并保留不同亮度和轻微角度变化的清晰样本。')
bullet(doc, '点击“清空未标注样本”可以删除当前模板所有历史批次中的未标注图片；数字标签和“不可辨认”样本都会保留。')

heading(doc, '选择模板参与识别', 2)
reset_numbers()
for text in [
    '进入“测量模板选择”页。',
    '为 1# 和 2# 分别勾选一个或多个模板；同一模板可以同时给两台使用。',
    '点击“保存模板选择并返回主界面”，或在视频相同时把当前采集框一并带回主界面。',
    '回到主界面后再次测试当前帧，确认模板可以识别，再开始完整分析。',
]: numbered(doc, text)

heading(doc, '分享样本模板', 2)
body(doc, '点击“导出当前模板包”会生成 ZIP，只包含已经确认的数字图片。收到模板包的人在第三页点击“导入别人发来的模板包”，导入后勾选即可使用。发送单文件 EXE 不会自动包含你后来建立的样本模板。')

heading(doc, '七 核对结果和保存到存档', 1)
body(doc, '分析结束后不要立刻导出，先核对目标点表格。容量正确但时间显示“--”或“时间待核对”时，可以保留容量并手工补时间；软件不会拿视频播放时间冒充仪表时间。')
reset_numbers()
for text in [
    '点击某一行的视频时刻，回看对应画面。',
    '双击时间或容量单元格修改错误值。',
    '确认 1#、2# 的充放电类型和 Excel 测试项没有选错。',
    '点击“保存本次结果到存档”。已有测试项重做时，软件会询问是否替换。',
    '换下一段视频，继续保存到同一个存档，直到计划项目完成。',
]: numbered(doc, text)

heading(doc, '打开工作版 Excel 插图', 2)
body(doc, '点击“打开存档 Excel”可以检查当前累计结果并手工插入产品图片。编辑完成后先在 WPS 或 Excel 中保存并关闭，再回到软件继续写入或最终导出。')
body(doc, '空白测试块出现 DIV 0 通常表示对应区域还没有数据。最终导出只保留已经完成的测试块。', '空白测试块出现 DIV 0')

heading(doc, '八 导出和发送最终结果', 1)
bullet(doc, '“导出存档完整 Excel 报告”会汇总该存档所有已完成项目，并保留相关公式和已插入图片。')
bullet(doc, '“导出完整数据为 CSV”适合后续统计或交给其他程序处理。')
bullet(doc, '最终 Excel 应另存到明确位置，不要覆盖存档工作版 report.xlsx。')
bullet(doc, '把软件发给别人时，只发送 PowerbankTestTool_v2_单文件版.exe 即可；如需共享数码管识别能力，再附上导出的 ZIP 模板包。')

heading(doc, '九 常见问题处理', 1)
table(doc, ['现象', '处理方法'], [
    ('第一次启动较慢', '单文件版正在解压 OCR 和界面组件，等待主窗口出现；后续启动通常会快一些。'),
    ('鼠标滚轮不能改参数', '这是防误触设计。点击输入框后键盘输入，或使用右侧上下箭头。'),
    ('存档无法写入或删除', '先保存并关闭正在打开的工作版 Excel，再重试。'),
    ('数码管大量未识别', '确认已选择数码管模式并勾选正确模板；缩小电量框，补齐缺失电量的清晰样本。'),
    ('数字闪烁或缺笔画', '采集间隔缩短至 0.2 秒，坏帧标为不可辨认，不要猜标签。'),
    ('识别结果时间偏后', '数码管模式会保存第一次候选画面的时间和容量；继续减小抽帧间隔可缩小定位误差。'),
    ('时间显示为待核对', 'Power Z 时间 OCR 未能可靠读取。点击视频时刻回看，双击表格手工补入。'),
    ('切换存档被阻止', '先结束视频分析；有未保存表格结果时，先保存或确认放弃。'),
    ('别人导入 EXE 后没有模板', '用户新建的样本库不在 EXE 内，需要单独导出 ZIP 模板包并让对方导入。'),
], widths=[4.2, 11.8])

heading(doc, '十 交付前检查清单', 1)
for text in [
    '所有视频结果都已保存到正确存档。',
    '1# 和 2# 的充放电类型、测试项对应正确。',
    '时间待核对项目已经回看并处理。',
    '工作版 Excel 中需要的产品图片已经插入并保存。',
    'WPS 或 Excel 已关闭。',
    '最终完整 Excel 已导出到单独位置并打开检查。',
    '需要共享的数码管模板已经导出为 ZIP。',
]: bullet(doc, '□ ' + text)

body(doc, '本地数据位置：存档和样本库默认保存在 %LOCALAPPDATA%\\PowerbankTestTool。软件不会把采集图片或标注自动上传网络。')

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = '充电宝数显均匀度测试工具使用攻略'
doc.core_properties.subject = 'PowerbankTestTool v2 操作手册'
doc.core_properties.author = 'PowerbankTestTool'
doc.save(OUTPUT)
print(OUTPUT)
