"""Real PB313 UI/worker acceptance probe. Requires the user's 20261012 video.

Runs against isolated settings, snapshots and archives under --output. Never
modifies the source video, sample library or the user's existing projects.
The expected ROIs and readings below belong to this particular video.
"""
import argparse
import os,sys,time,json
from datetime import datetime
from pathlib import Path
BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--video', type=Path, required=True)
parser.add_argument('--library', type=Path, required=True)
parser.add_argument('--output', type=Path, default=BASE/'测试结果'/('capture-'+datetime.now().strftime('%Y%m%d-%H%M%S')))
parser.add_argument('--show-ui', action='store_true')
args = parser.parse_args()
source, LIB, OUT = args.video.resolve(), args.library.resolve(), args.output.resolve()
if not source.is_file() or not LIB.is_dir():
    parser.error('视频或样本库路径不存在')
OUT.mkdir(parents=True, exist_ok=False)
os.environ['LOCALAPPDATA'] = str(OUT/'appdata')
os.environ['QT_QPA_PLATFORM'] = 'windows' if args.show_ui else 'offscreen'
import onnxruntime,cv2
from PyQt5.QtCore import QSettings,QTimer,Qt
from PyQt5.QtWidgets import QApplication,QMessageBox,QFileDialog
from PyQt5.QtTest import QTest
from unittest.mock import patch
from core.segment_dataset import read_frame_at,SEGMENT_MODE,read_image
from core.project_archive import create_project,ProjectArchive,group_key
from core.record_evidence import load_snapshot
from core.video_worker import VideoAnalysisWorker
from ui.main_window import MainWindow
from ui.manual_capture_dialog import ManualCaptureDialog
import numpy as np
QSettings.setDefaultFormat(QSettings.IniFormat)
QSettings.setPath(QSettings.IniFormat,QSettings.UserScope,str(OUT/'settings'))
app=QApplication([])
from PyQt5.QtGui import QFont
app.setFont(QFont('Microsoft YaHei',9))
app.setStyle('Fusion')
rois={1:{'soc':[949,197,73,47],'time':[169,341,60,17],'cap':[230,341,50,17]},
      2:{'soc':[1287,252,76,49],'time':[159,769,58,17],'cap':[218,769,50,17]}}
results={'source':str(source),'library':str(LIB),'manual':{},'manual_missing':{},'continuous':{}}

def wait_task(dialog):
    deadline=time.monotonic()+90
    while dialog.task and dialog.task.isRunning():
        app.processEvents();time.sleep(.01)
        if time.monotonic()>deadline:
            dialog.task.requestInterruption();dialog.task.wait()
            raise AssertionError('worker timeout')
    for _ in range(3):app.processEvents()

project=create_project('PB313 real-frame manual capture validation',next(BASE.glob('PB519*.xlsx')))
with patch('ui.main_window.ResourceMonitor.start'):
    window=MainWindow()
window.reload_projects(str(project.root))
window.video_path=str(source)
window.video_cap=cv2.VideoCapture(str(source))
window.fps=window.video_cap.get(cv2.CAP_PROP_FPS)
window.total_frames=int(window.video_cap.get(cv2.CAP_PROP_FRAME_COUNT))
ok,window.current_frame_cv=read_frame_at(window.video_cap,11500)
assert ok
for unit in (1,2):
    window.unit_modes[unit].setCurrentIndex(window.unit_modes[unit].findData('charge'))
    for name,roi in rois[unit].items():window.canvas.rois[f'{name}_{unit}']=roi
window.show();app.processEvents()
for unit in (1, 2):
    assert getattr(window, f'manual_{unit}').isVisible()
    assert getattr(window, f'visual_{unit}').isVisible()
for unit, target_soc, selected_second, expected in [(1,100,11500,('03:11:33',8.5088)), (2,100,11500,('03:11:32',8.4370)), (1,99,11450,('03:10:43',8.5031)), (2,99,11150,('03:05:42',8.4068))]:
    errors=[]
    def drive():
        dialog=app.activeModalWidget()
        try:
            assert isinstance(dialog,ManualCaptureDialog)
            assert dialog.unit==unit and dialog.soc==target_soc
            dialog.start_spin.setValue(selected_second-.2);dialog.end_spin.setValue(selected_second+.2);dialog.step_spin.setValue(.1)
            QTest.mouseClick(dialog.scan_button,Qt.LeftButton);wait_task(dialog)
            assert dialog.candidates.count()>=3,dialog.status.text()
            closest=min(range(dialog.candidates.count()),key=lambda i:abs(dialog.candidates.item(i).data(Qt.UserRole)['video_sec']-selected_second))
            dialog.candidates.setCurrentRow(closest);app.processEvents()
            assert abs(dialog.selected['video_sec']-selected_second)<.04
            old_index=dialog.selected['frame_index']
            QTest.mouseClick(dialog.next_button,Qt.LeftButton);wait_task(dialog)
            assert dialog.selected['frame_index']==old_index+1,(dialog.selected,old_index)
            QTest.mouseClick(dialog.previous_button,Qt.LeftButton);wait_task(dialog)
            assert dialog.selected['frame_index']==old_index
            dialog.single_line.setChecked(True)
            QTest.mouseClick(dialog.read_button,Qt.LeftButton);wait_task(dialog)
            assert dialog.fields is not None,dialog.status.text()
            assert dialog.time_edit.text()==expected[0],dialog.fields
            assert abs(float(dialog.capacity_edit.text())-expected[1])<.00005,dialog.fields
            dialog.grab().save(str(OUT/f'manual-unit-{unit}-{target_soc}.png'))
            QTest.mouseClick(dialog.confirm_button,Qt.LeftButton)
        except BaseException as exc:
            import traceback
            errors.append(traceback.format_exc());dialog.reject()
    table=window.table_1 if unit==1 else window.table_2
    from core.excel_filler import get_target_soc_list
    table.setCurrentCell(get_target_soc_list('charge').index(target_soc),0)
    QTimer.singleShot(50,drive)
    QTest.mouseClick(getattr(window,f'manual_{unit}'),Qt.LeftButton)
    if errors:raise AssertionError(errors)
    record=(window.extracted_records_1 if unit==1 else window.extracted_records_2)[target_soc]
    results['manual' if target_soc==100 else 'manual_missing'][str(unit)]={k:v for k,v in record.items() if k!='time_obj'}
    check=cv2.VideoCapture(str(source))
    ok,exact=read_frame_at(check,record['video_sec']);check.release()
    assert ok and np.array_equal(load_snapshot(record),exact)
    print('MANUAL_OK',unit,target_soc,record['time_str'],record['capacity'],record['frame_index'],flush=True)
window.grab().save(str(OUT/'main-dual-records.png'))
with patch.object(QMessageBox,'information'),patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):
    window.save_to_project()
assert not window.dirty_units
reopened=ProjectArchive(project.root)
assert len(reopened.metadata['entries'])==2
for entry in reopened.metadata['entries'].values():
    for soc in ('99','100'):
        rec=entry['records'][soc]
        assert rec['capture_method']=='manual' and rec['capture_unit']==entry['unit']
        assert reopened.resolve_record(rec)['snapshot_path']
output=OUT/'manual-dual-export.xlsx'
with patch.object(QFileDialog,'getSaveFileName',return_value=(str(output),'Excel')),patch.object(QMessageBox,'information'):
    window.export_to_excel()
assert output.exists()
import openpyxl
wb=openpyxl.load_workbook(output,data_only=False)
from core.excel_filler import get_cell_coords
for unit in (1,2):
    group=window.unit_groups[unit].currentData()
    for soc,key in [(99,'manual_missing'),(100,'manual')]:
        record=results[key][str(unit)]
        coordinates=get_cell_coords(soc,'charge',group['title_row']-2,group['col_offset'])
        from datetime import time as clock_time
        assert wb[group['sheet']].cell(*coordinates[0]).value==clock_time.fromisoformat(record['time_str'])
        assert abs(wb[group['sheet']].cell(*coordinates[1]).value-record['capacity'])<.00005
        print('EXPORT_OK',unit,soc,coordinates,flush=True)
wb.close()
results['archive']=str(project.root);results['export']=str(output)
window.close();app.processEvents()

for high_speed in (False,True):
    worker=VideoAnalysisWorker(str(source),**{f'roi_{k}_{u}':v for u,r in rois.items() for k,v in r.items()},
        mode='charge',active_units=[1,2],sample_interval_sec=.5,debounce_count=2,
        digit_type=SEGMENT_MODE,segment_paths={1:str(LIB),2:str(LIB)},
        visual_tracking={1:True,2:True},meter_single_line={1:True,2:True},high_speed=high_speed,cpu_threads=2,
        visual_reference_records={u:[dict(results['manual'][str(u)],soc=100),dict(results['manual_missing'][str(u)],soc=99)] for u in (1,2)},
        start_sec=11464,end_sec=11476)
    errors=[];records=[]
    worker.error_signal.connect(errors.append)
    worker.progress_signal.connect(lambda sec,*args: print('AT',high_speed,round(sec,2),flush=True))
    worker.finished_signal.connect(lambda a,b:records.append((a,b)))
    worker.run()
    assert not errors,errors
    assert records
    summary={str(u):[{k:v for k,v in rec.items() if k not in ('time_obj','snapshot_path')}
                     for rec in records[0][u-1].values()] for u in (1,2)}
    results['continuous'][str(high_speed)]=summary
    print('CONTINUOUS',high_speed,json.dumps(summary,ensure_ascii=False),flush=True)
    for unit in (1, 2):
        for soc, rec in records[0][unit-1].items():
            snapshot = load_snapshot(rec)
            assert snapshot is not None
            from core.segment_dataset import write_image
            write_image(OUT/f'continuous-{high_speed}-{unit}-{soc}.png', snapshot)
            assert rec['time_obj'] is not None and rec['capacity'] is not None, rec
            # Re-read the saved frame through the manual route and require identical fields.
            from core.manual_capture import selected_fields
            fields = selected_fields(snapshot, rois[unit], {}, worker.ocr_engine, single_line=True)
            assert fields['time_str'] == rec['time_str'] and fields['capacity'] == rec['capacity']

(OUT/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
assert results['continuous']['False']==results['continuous']['True']
assert [r['soc'] for r in results['continuous']['False']['1']]==[99,100], 'Real transition not captured'
assert [r['soc'] for r in results['continuous']['False']['2']]==[100], 'Second unit constant value not captured'
(OUT/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print('VALIDATION_DONE',OUT,flush=True)
