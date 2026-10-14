import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import onnxruntime
import tempfile
import unittest
from datetime import time
from pathlib import Path
from unittest.mock import patch

import cv2
import numpy as np
import openpyxl
from openpyxl.drawing.image import Image
from PyQt5.QtWidgets import QApplication

from core.project_archive import ProjectArchive, create_project, delete_project, group_key


class ProjectArchiveTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_mixed_charge_discharge_reopen_and_preserve_inserted_picture(self):
        template = next(Path(__file__).resolve().parents[1].glob('PB519*.xlsx'))
        with tempfile.TemporaryDirectory() as temporary, patch(
                'core.project_archive.projects_home', return_value=Path(temporary) / 'projects'):
            project = create_project('PB519', template)
            self.assertEqual(openpyxl.load_workbook(project.report_path)['充电时间及曲线']['AA5'].value,
                             '=Z5/$Y$32')
            charge = next(g for g in project.groups['充电时间及曲线'] if g['col_offset'] == 0)
            discharge = next(g for g in project.groups['放电时间及曲线']
                             if g['col_offset'] == 14 and g['title_row'] == 37)
            project.set_planned([group_key(charge), group_key(discharge)])
            project.save_results([{'unit': 1, 'mode': 'charge', 'group': charge,
                'records': {1: {'time_str': '00:00:07', 'time_obj': time(0, 0, 7),
                                'capacity': 0.01, 'video_sec': 3.5}}}], 'first.mp4')
            picture = Path(temporary) / 'picture.png'
            cv2.imwrite(str(picture), np.full((20, 20, 3), 150, np.uint8))
            workbook = openpyxl.load_workbook(project.report_path)
            image = Image(str(picture))
            workbook['充电时间及曲线'].add_image(image, 'B10')
            workbook.save(project.report_path)
            project = ProjectArchive(project.root)
            project.save_results([{'unit': 2, 'mode': 'discharge', 'group': discharge,
                'records': {99: {'time_str': '00:01:00', 'time_obj': time(0, 1),
                                 'capacity': 0.08, 'video_sec': 31.0}}}], 'second.mp4')
            self.assertEqual(len(project.metadata['entries']), 2)
            self.assertEqual(len(openpyxl.load_workbook(project.report_path)['充电时间及曲线']._images), 1)
            final = Path(temporary) / 'final.xlsx'
            self.assertEqual(project.export(final), 2)
            result = openpyxl.load_workbook(final)
            self.assertEqual(set(result.sheetnames), {'充电时间及曲线', '放电时间及曲线'})
            self.assertEqual(result['充电时间及曲线']['B4'].value, time(0, 0, 7))
            self.assertEqual(result['放电时间及曲线']['P4'].value, time(0, 1))
            self.assertEqual(len(result['充电时间及曲线']._images), 1)
            self.assertLessEqual(result['充电时间及曲线'].max_row, 35)
            self.assertIsNone(result['充电时间及曲线']['O2'].value)

    def test_export_moves_later_block_formulas_and_images(self):
        template = next(Path(__file__).resolve().parents[1].glob('PB519*.xlsx'))
        with tempfile.TemporaryDirectory() as temporary, patch(
                'core.project_archive.projects_home', return_value=Path(temporary) / 'projects'):
            project = create_project('PB519', template)
            first = next(g for g in project.groups['充电时间及曲线']
                         if g['title_row'] == 2 and g['col_offset'] == 0)
            third = next(g for g in project.groups['充电时间及曲线']
                         if g['title_row'] == 72 and g['col_offset'] == 0)
            for group in (first, third):
                project.save_results([{'unit': 1, 'mode': 'charge', 'group': group,
                    'records': {1: {'time_obj': time(0, 0, 7), 'capacity': 0.02}}}], 'video.mp4')
            picture = Path(temporary) / 'later.png'
            cv2.imwrite(str(picture), np.full((20, 20, 3), 150, np.uint8))
            workbook = openpyxl.load_workbook(project.report_path)
            workbook['充电时间及曲线'].add_image(Image(str(picture)), 'B80')
            workbook.save(project.report_path)
            final = Path(temporary) / 'final.xlsx'
            project.export(final)
            result = openpyxl.load_workbook(final)
            ws = result['充电时间及曲线']
            self.assertLessEqual(ws.max_row, 70)
            self.assertGreaterEqual(ws.max_row, 69)
            self.assertIn('12V', ws['A37'].value)
            self.assertEqual(ws['M39'].value, '=L39/K66')
            self.assertEqual(ws._images[0].anchor._from.row, 44)

    def test_two_units_share_a_voltage_block_without_overwriting(self):
        template = next(Path(__file__).resolve().parents[1].glob('PB519*.xlsx'))
        with tempfile.TemporaryDirectory() as temporary, patch(
                'core.project_archive.projects_home', return_value=Path(temporary) / 'projects'):
            project = create_project('PB519', template)
            pair = [g for g in project.groups['放电时间及曲线'] if g['title_row'] == 2]
            assignments = [{'unit': 1 if g['col_offset'] == 0 else 2,
                'mode': 'discharge', 'group': g,
                'records': {99: {'time_obj': time(0, 0, 7 if g['col_offset'] == 0 else 9),
                                 'capacity': 0.05}}} for g in pair]
            project.save_results(assignments, 'both.mp4')
            final = Path(temporary) / 'both.xlsx'
            project.export(final)
            report = openpyxl.load_workbook(final)
            ws = report['放电时间及曲线']
            self.assertEqual(ws['B4'].value, time(0, 0, 7))
            self.assertEqual(ws['P4'].value, time(0, 0, 9))
            self.assertIn('1#', ws['A2'].value)
            self.assertIn('2#', ws['O2'].value)

    def test_synthetic_dual_unit_time_capacity_round_trip_export(self):
        """Synthetic fixture values must remain in their unit's cells and archive rows."""
        template = next(Path(__file__).resolve().parents[1].glob('PB519*.xlsx'))
        with tempfile.TemporaryDirectory() as temporary, patch(
                'core.project_archive.projects_home', return_value=Path(temporary) / 'projects'):
            project = create_project('synthetic dual-unit export', template)
            discharge_sheet = next(name for name in project.groups if '\u653e\u7535' in name)
            groups = [g for g in project.groups[discharge_sheet] if g['title_row'] == 2]
            group_by_unit = {1 if g['col_offset'] == 0 else 2: g for g in groups}
            fixture = {
                1: {'time_str': '00:12:34', 'time_obj': time(0, 12, 34), 'capacity': 1.2345},
                2: {'time_str': '00:56:07', 'time_obj': time(0, 56, 7), 'capacity': 0.0678},
            }
            assignments = [
                {'unit': unit, 'mode': 'discharge', 'group': group_by_unit[unit],
                 'records': {99: dict(record, video_sec=unit * 10.0)}}
                for unit, record in fixture.items()
            ]
            self.assertEqual(project.save_results(assignments, 'synthetic-fixture-video.mp4'),
                             {1: 1, 2: 1})

            reopened = ProjectArchive(project.root)
            output = Path(temporary) / 'synthetic-export.xlsx'
            self.assertEqual(reopened.export(output), 2)
            result = openpyxl.load_workbook(output, data_only=False)
            ws = result[discharge_sheet]
            self.assertEqual(ws['B4'].value, time(0, 12, 34))
            self.assertEqual(ws['P4'].value, time(0, 56, 7))
            self.assertEqual(ws['B5'].value, 1.2345)
            self.assertEqual(ws['P5'].value, 0.0678)
            self.assertEqual(ws['B4'].number_format, 'h:mm:ss')
            self.assertEqual(ws['P4'].number_format, 'h:mm:ss')
            self.assertEqual(ws['B5'].number_format, '0.0000_ ')
            self.assertEqual(ws['P5'].number_format, '0.0000_ ')
            for unit, group in group_by_unit.items():
                entry = reopened.metadata['entries'][group_key(group)]
                self.assertEqual(entry['unit'], unit)
                self.assertEqual(entry['records']['99']['time_str'], fixture[unit]['time_str'])
                self.assertEqual(entry['records']['99']['capacity'], fixture[unit]['capacity'])
            result.close()

    def test_delete_project_is_scoped_and_rejects_open_excel(self):
        template = next(Path(__file__).resolve().parents[1].glob('PB519*.xlsx'))
        with tempfile.TemporaryDirectory() as temporary, patch(
                'core.project_archive.projects_home', return_value=Path(temporary) / 'projects'):
            project = create_project('delete me', template)
            lock = project.root / '~$report.xlsx'
            lock.touch()
            with self.assertRaises(PermissionError):
                delete_project(project.root)
            self.assertTrue(project.root.exists())
            lock.unlink()
            delete_project(project.root)
            self.assertFalse(project.root.exists())
            outside = Path(temporary) / 'outside'
            outside.mkdir()
            (outside / 'project.json').write_text('{}', encoding='utf-8')
            (outside / 'report.xlsx').touch()
            with self.assertRaises(ValueError):
                delete_project(outside)

    def test_window_routes_each_unit_to_its_own_sheet_and_group(self):
        from ui.main_window import MainWindow
        template = next(Path(__file__).resolve().parents[1].glob('PB519*.xlsx'))
        with tempfile.TemporaryDirectory() as temporary, patch(
                'core.project_archive.projects_home', return_value=Path(temporary) / 'projects'):
            project = create_project('PB519', template)
            with patch('ui.main_window.ResourceMonitor.start'):
                window = MainWindow()
            window.reload_projects(str(project.root))
            window.unit_modes[1].setCurrentIndex(1)
            self.assertEqual(window.unit_groups[1].currentData()['sheet'], '充电时间及曲线')
            self.assertEqual(window.unit_groups[2].currentData()['sheet'], '放电时间及曲线')
            window.on_worker_point_detected(1, {'soc': 1, 'time_str': '00:00:07',
                'time_obj': time(0, 0, 7), 'capacity': 0.01, 'video_sec': 3.5})
            window.on_worker_point_detected(2, {'soc': 99, 'time_str': '00:01:00',
                'time_obj': time(0, 1), 'capacity': 0.08, 'video_sec': 31.0})
            self.assertEqual(window.table_1.item(0, 0).text(), '1%')
            self.assertEqual(window.table_2.item(0, 0).text(), '99%')
            with patch('ui.main_window.QMessageBox.information'):
                window.save_to_project()
            self.assertEqual(len(window.project.metadata['entries']), 2)
            self.assertEqual(window.dirty_units, set())
            window.close()
